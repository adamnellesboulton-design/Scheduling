"""Excel workbook output (Section 9), built with openpyxl.

Four sheets: Schedule (master grid), Summary, Compliance, Config.
Filename: dialysis_schedule_<start>_<end>.xlsx
"""

from __future__ import annotations

from datetime import timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from .config import Config
from .model import build_operating_dates, is_worked

# Clean black-and-white output: no decorative fills, so the workbook prints
# cleanly in B/W. Colour is reserved to FLAG ISSUES only (coverage shortfalls
# and FAIL/WARN compliance rows).
FILL_D10 = None
FILL_D5 = None
FILL_LV = None
FILL_SAT_COL = None
FILL_HEADER = None
FILL_PASS = None
FILL_INFO = None
FILL_SHORT = PatternFill("solid", fgColor="FF9999")  # coverage shortfall (issue)
FILL_FAIL = PatternFill("solid", fgColor="FFC7CE")  # FAIL (issue)
FILL_WARN = PatternFill("solid", fgColor="FFE08A")  # WARN (issue)

WHITE_BOLD = Font(bold=True)  # headers: bold black text on no fill (B/W safe)
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def output_filename(cfg: Config) -> str:
    start = cfg.start
    end = start + timedelta(weeks=cfg.weeks) - timedelta(days=1)
    return f"dialysis_schedule_{start.isoformat()}_{end.isoformat()}.xlsx"


def _print_setup(ws, subtitle: str, landscape: bool = True, period: str = "",
                 option: str = ""):
    """Landscape, fit-to-width, with a printed title header/footer.

    The header/footer print at the top/bottom of every page without shifting any
    cells, so the grid layout is untouched.
    """
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    right = " | ".join(p for p in (period, option) if p)
    ws.oddHeader.left.text = "BC Children's Hospital — Pediatric Dialysis Unit"
    ws.oddHeader.center.text = subtitle
    ws.oddHeader.right.text = right
    ws.oddFooter.left.text = "Generated &D"
    ws.oddFooter.right.text = "Page &P of &N"


def _style_cell(cell, fill=None, font=None, align=CENTER, border=True):
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    cell.alignment = align
    if border:
        cell.border = BORDER


def build_workbook(cfg: Config, result, report) -> Workbook:
    wb = Workbook()
    start = cfg.start
    end = start + timedelta(weeks=cfg.weeks) - timedelta(days=1)
    period = f"{start.strftime('%d %b %Y')} – {end.strftime('%d %b %Y')}"
    option = getattr(result, "label", "") or ""

    _build_schedule_sheet(wb, cfg, result)
    _build_summary_sheet(wb, report)
    _build_compliance_sheet(wb, report)
    _build_config_sheet(wb, cfg, result)

    _print_setup(wb["Schedule"], "Master Schedule", True, period, option)
    _print_setup(wb["Summary"], "Per-nurse Summary", False, period, option)
    _print_setup(wb["Compliance"], "Compliance Report", False, period, option)
    _print_setup(wb["Config"], "Configuration Snapshot", False, period, option)
    return wb


# --- Sheet 1: Schedule -----------------------------------------------------


def _build_schedule_sheet(wb: Workbook, cfg: Config, result):
    ws = wb.active
    ws.title = "Schedule"
    operating = result.operating or build_operating_dates(cfg)
    assignments = result.assignments

    # Header rows: row1 week numbers, row2 date labels.
    ws.cell(row=1, column=1, value="Nurse")
    ws.cell(row=2, column=1, value="")
    _style_cell(ws.cell(1, 1), FILL_HEADER, WHITE_BOLD)
    _style_cell(ws.cell(2, 1), FILL_HEADER, WHITE_BOLD)

    col_of_iso: dict[str, int] = {}
    for idx, od in enumerate(operating):
        col = 2 + idx
        col_of_iso[od.iso] = col
        wk_cell = ws.cell(1, col, value=f"Week {od.week_index + 1}")
        dt_cell = ws.cell(2, col, value=od.d.strftime("%a %d-%b"))
        sat_fill = FILL_SAT_COL if od.is_saturday else FILL_HEADER
        _style_cell(wk_cell, FILL_HEADER, WHITE_BOLD)
        _style_cell(dt_cell, sat_fill, BOLD if od.is_saturday else WHITE_BOLD)
        ws.column_dimensions[get_column_letter(col)].width = 11

    ws.column_dimensions["A"].width = 18

    # Per-nurse rows.
    first_data_row = 3
    n_cols = 1 + len(operating)
    sat_cols = {col for iso, col in col_of_iso.items()
                if any(od.iso == iso and od.is_saturday for od in operating)}

    for r, nurse in enumerate(cfg.nurses):
        row = first_data_row + r
        name_cell = ws.cell(row, 1, value=nurse.name)
        _style_cell(name_cell, font=BOLD, align=Alignment(horizontal="left"))
        worked = assignments.get(nurse.name, {})
        unavailable = set(nurse.unavailable_dates)
        for od in operating:
            col = col_of_iso[od.iso]
            code = worked.get(od.iso)
            if code:
                fill = FILL_D10 if code == "D10" else FILL_D5
                val = code
            elif od.iso in unavailable:
                fill = FILL_LV
                val = "LV"
            elif od.is_saturday:
                fill = FILL_SAT_COL
                val = ""
            else:
                fill = None
                val = ""
            cell = ws.cell(row, col, value=val)
            _style_cell(cell, fill)

    # Right-hand per-nurse columns: total hours, FTE, Saturdays.
    from .fte import scheduled_fte

    od_by_iso = {od.iso: od for od in operating}
    summary_start = n_cols + 1
    headers = ["Total hrs", "Sched FTE", "Sats worked"]
    for j, h in enumerate(headers):
        c = ws.cell(2, summary_start + j, value=h)
        _style_cell(c, FILL_HEADER, WHITE_BOLD)
        ws.column_dimensions[get_column_letter(summary_start + j)].width = 11
    for r, nurse in enumerate(cfg.nurses):
        row = first_data_row + r
        worked = assignments.get(nurse.name, {})
        total = sum(od_by_iso[i].paid_hours for i, c in worked.items()
                    if i in od_by_iso and is_worked(c))
        sats = sum(1 for i, c in worked.items()
                   if i in od_by_iso and od_by_iso[i].is_saturday and is_worked(c))
        _style_cell(ws.cell(row, summary_start, value=round(total, 1)))
        _style_cell(
            ws.cell(row, summary_start + 1, value=round(scheduled_fte(total, cfg.weeks), 3))
        )
        _style_cell(ws.cell(row, summary_start + 2, value=sats))

    # Bottom per-day assigned vs required row.
    count_row = first_data_row + len(cfg.nurses)
    lab = ws.cell(count_row, 1, value="Assigned / Required")
    _style_cell(lab, FILL_HEADER, WHITE_BOLD, Alignment(horizontal="left"))
    for od in operating:
        col = col_of_iso[od.iso]
        assigned = sum(1 for name in assignments
                       if is_worked(assignments[name].get(od.iso)))
        cell = ws.cell(count_row, col, value=f"{assigned}/{od.demand}")
        fill = FILL_SHORT if assigned < od.demand else None
        _style_cell(cell, fill, BOLD if assigned < od.demand else None)

    # Freeze panes: names in col A, two header rows.
    ws.freeze_panes = "B3"

    # Legend block beneath the grid (Section 9).
    legend_row = count_row + 2
    legend = [
        ("LEGEND", BOLD),
        ("D10 = 0730-1730 (10.0h elapsed, 9.5h paid with the 30-min unpaid "
         "meal). A missed meal is paid as overtime (Art. 27, flagged not priced).",
         None),
        ("D5 = 0730-1230 (5.0h, paid). No meal period required (26.03(A) "
         "triggers only beyond 5 consecutive hours).", None),
        ("Blank = off.  LV = unavailable/approved leave.  ST = statutory holiday "
         "(paid, not worked) — scheduled on the actual holiday date.", None),
        ("Meal/rest (26.03 / 26.04): D10 30-min meal must begin by 1230 "
         "(<=5.0h after start); two paid 15-min rests per D10. D5: one paid "
         "15-min rest (shift >= 4h).", None),
        ("Extended Work Day (25.11 / 26.01): D10 exceeds the 7.5h normal daily "
         "full shift -- verify against your Extended Work Day Memorandum terms.",
         None),
        ("Cells are shaded only to flag problems (red = coverage shortfall); the "
         "grid is otherwise plain for clean black-and-white printing.", None),
    ]
    for i, (text, font) in enumerate(legend):
        c = ws.cell(legend_row + i, 1, value=text)
        if font:
            c.font = font
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)


# --- Sheet 2: Summary ------------------------------------------------------


def _build_summary_sheet(wb: Workbook, report):
    ws = wb.create_sheet("Summary")
    headers = [
        "Nurse", "D10 (sched/target)", "D5 (sched/target)", "Counts met",
        "Sched FTE", "Total hrs", "Avg hrs/wk", "Sats worked", "Sats in period",
        "Worst 9-wk Sat",
    ]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(1, j, value=h)
        _style_cell(c, FILL_HEADER, WHITE_BOLD)
        ws.column_dimensions[get_column_letter(j)].width = 16
    for r, s in enumerate(report.nurse_summaries, start=2):
        counts_met = (s.scheduled_d10 == s.target_d10
                      and s.scheduled_d5 == s.target_d5)
        vals = [
            s.name, f"{s.scheduled_d10}/{s.target_d10}",
            f"{s.scheduled_d5}/{s.target_d5}", "YES" if counts_met else "NO",
            s.scheduled_fte, s.total_hours, s.avg_weekly_hours,
            s.saturdays_worked, s.saturdays_in_period, s.worst_9wk_sat,
        ]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(r, j, value=v)
            fill = None
            if j == 4 and not counts_met:  # flag only the issue
                fill = FILL_WARN
            _style_cell(c, fill, align=Alignment(
                horizontal="left" if j == 1 else "center"))
    ws.freeze_panes = "A2"


# --- Sheet 3: Compliance ---------------------------------------------------


def _build_compliance_sheet(wb: Workbook, report):
    ws = wb.create_sheet("Compliance")
    headers = ["Rule", "Article", "Status", "Detail"]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(1, j, value=h)
        _style_cell(c, FILL_HEADER, WHITE_BOLD)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 80
    for r, rule in enumerate(report.rules, start=2):
        # Only colour the problems (FAIL/WARN); PASS/INFO stay plain for B/W.
        fill = {"FAIL": FILL_FAIL, "WARN": FILL_WARN}.get(rule.status)
        _style_cell(ws.cell(r, 1, value=rule.rule), align=Alignment(
            horizontal="left", wrap_text=True, vertical="top"))
        _style_cell(ws.cell(r, 2, value=rule.citation), align=Alignment(
            horizontal="left", vertical="top"))
        _style_cell(ws.cell(r, 3, value=rule.status), fill, BOLD)
        _style_cell(ws.cell(r, 4, value=rule.detail), align=Alignment(
            horizontal="left", wrap_text=True, vertical="top"))
    ws.freeze_panes = "A2"


# --- Sheet 4: Config -------------------------------------------------------


def _build_config_sheet(wb: Workbook, cfg: Config, result):
    ws = wb.create_sheet("Config")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 60

    def kv(row, k, v):
        kc = ws.cell(row, 1, value=k)
        kc.font = BOLD
        ws.cell(row, 2, value=v)
        return row + 1

    r = 1
    c = ws.cell(r, 1, value="Schedule generation snapshot (audit record, 25.05)")
    c.font = BOLD
    r += 2
    r = kv(r, "Start date", cfg.start_date)
    r = kv(r, "Weeks in period", cfg.weeks)
    end = cfg.start + timedelta(weeks=cfg.weeks) - timedelta(days=1)
    r = kv(r, "End date", end.isoformat())
    r = kv(r, "Demand (Mon/Wed/Fri/Sat)",
           ", ".join(f"{k}={v}" for k, v in cfg.demand.items()))
    if cfg.weekly_demand_override:
        r = kv(r, "Weekly demand overrides", str(cfg.weekly_demand_override))
    r = kv(r, "Meal handling", "30-min unpaid meal (D10 = 9.5h paid); "
                               "missed meals paid as OT (Art. 27)")
    r = kv(r, "Default FTE flex", cfg.fte_tolerance)
    r = kv(r, "Weekly full-time hours", cfg.weekly_full_time_hours)
    r = kv(r, "Generation method", result.method)
    r = kv(r, "Solver status", result.status)
    r = kv(r, "Extra flex applied", result.tolerance_used)
    r += 1

    c = ws.cell(r, 1, value="Operating shifts")
    c.font = BOLD
    r += 1
    for s in cfg.operating_shifts:
        paid = s.paid_hours(cfg.meal_designated_available)
        r = kv(
            r,
            f"  {s.weekday_name} {s.code}",
            f"{s.start}-{s.end}, {s.elapsed_hours}h elapsed, {paid}h paid",
        )
    r += 1

    c = ws.cell(r, 1, value="Roster")
    c.font = BOLD
    r += 1
    hdr = ["Name", "D10", "D5", "Stat", "FTE (derived)", "Job share",
           "Preferences", "Unavailable dates"]
    for j, h in enumerate(hdr, start=1):
        cc = ws.cell(r, j, value=h)
        _style_cell(cc, FILL_HEADER, WHITE_BOLD)
    r += 1
    for nurse in cfg.nurses:
        prefs = []
        if nurse.pref_nonconsec_sat:
            prefs.append("non-consec Sat")
        if nurse.pref_clustered:
            prefs.append("cluster shifts")
        if nurse.fixed_off_mon:
            prefs.append("Mon off (fixed)")
        if nurse.fixed_work_weekly:
            prefs.append("work weekly")
        if nurse.fixed_fri_before_sat:
            prefs.append("Fri before Sat")
        if nurse.pref_off_wed:
            prefs.append("off Wed")
        if nurse.pref_off_fri:
            prefs.append("off Fri")
        ws.cell(r, 1, value=nurse.name)
        ws.cell(r, 2, value=nurse.target_d10)
        ws.cell(r, 3, value=nurse.target_d5)
        ws.cell(r, 4, value=nurse.stat_days)
        ws.cell(r, 5, value=nurse.target_fte)
        ws.cell(r, 6, value=nurse.job_share_group or "-")
        ws.cell(r, 7, value=", ".join(prefs) if prefs else "-")
        ws.cell(r, 8, value=", ".join(nurse.unavailable_dates))
        r += 1


def workbook_bytes(cfg: Config, result, report) -> bytes:
    wb = build_workbook(cfg, result, report)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
